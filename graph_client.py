import os
import io
import json
import logging
import aiohttp
import pdfplumber
import docx
import openpyxl
from azure.identity.aio import ClientSecretCredential
from dotenv import load_dotenv

load_dotenv()

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("graph_client")

TENANT_ID = os.getenv("TENANT_ID")
CLIENT_ID = os.getenv("MICROSOFT_APP_ID")
CLIENT_SECRET = os.getenv("MICROSOFT_APP_PASSWORD")
SHAREPOINT_HOSTNAME = os.getenv("SHAREPOINT_HOSTNAME")
SHAREPOINT_SITE_NAME = os.getenv("SHAREPOINT_SITE_NAME")
SHAREPOINT_FOLDER = os.getenv("SHAREPOINT_FOLDER_PATH", "")

GRAPH_BASE = "https://graph.microsoft.com/v1.0"

logger.info(f"SHAREPOINT_HOSTNAME={SHAREPOINT_HOSTNAME}")
logger.info(f"SHAREPOINT_SITE_NAME={SHAREPOINT_SITE_NAME}")
logger.info(f"SHAREPOINT_FOLDER={SHAREPOINT_FOLDER}")


async def get_token():
    credential = ClientSecretCredential(TENANT_ID, CLIENT_ID, CLIENT_SECRET)
    token = await credential.get_token("https://graph.microsoft.com/.default")
    await credential.close()
    return token.token


async def get_site_id():
    token = await get_token()
    url = f"{GRAPH_BASE}/sites/{SHAREPOINT_HOSTNAME}:/sites/{SHAREPOINT_SITE_NAME}"
    async with aiohttp.ClientSession() as session:
        async with session.get(url, headers={"Authorization": f"Bearer {token}"}) as resp:
            data = await resp.json()
            logger.info(f"get_site_id: {data.get('id', 'NOT FOUND')}")
            return data.get("id")


async def list_files(folder_path: str = ""):
    """מחזיר רשימת קבצים ותיקיות מתיקייה מסוימת"""
    token = await get_token()
    site_id = await get_site_id()

    target = folder_path or SHAREPOINT_FOLDER
    if target:
        url = f"{GRAPH_BASE}/sites/{site_id}/drive/root:/{target}:/children"
    else:
        url = f"{GRAPH_BASE}/sites/{site_id}/drive/root/children"

    logger.info(f"list_files URL: {url}")
    async with aiohttp.ClientSession() as session:
        async with session.get(url, headers={"Authorization": f"Bearer {token}"}) as resp:
            data = await resp.json()

    items = []
    for item in data.get("value", []):
        name = item.get("name", "")
        if item.get("folder"):
            items.append({
                "name": name,
                "id": item.get("id"),
                "type": "folder",
                "childCount": item.get("folder", {}).get("childCount", 0)
            })
        else:
            items.append({
                "name": name,
                "id": item.get("id"),
                "type": "file",
                "size": item.get("size"),
                "lastModified": item.get("lastModifiedDateTime"),
                "webUrl": item.get("webUrl", "")
            })
    logger.info(f"list_files found {len(items)} items")
    return items


async def list_folder(folder_name: str):
    """נכנס לתת-תיקייה וקורא את הקבצים שבתוכה"""
    target = SHAREPOINT_FOLDER
    if target:
        full_path = f"{target}/{folder_name}"
    else:
        full_path = folder_name
    return await list_files(folder_path=full_path)


async def search_files(query: str):
    """מחפש קבצים לפי מילות מפתח בכל ה-Drive"""
    token = await get_token()
    site_id = await get_site_id()
    url = f"{GRAPH_BASE}/sites/{site_id}/drive/root/search(q='{query}')"
    logger.info(f"search_files URL: {url}")
    async with aiohttp.ClientSession() as session:
        async with session.get(url, headers={"Authorization": f"Bearer {token}"}) as resp:
            data = await resp.json()

    files = []
    for item in data.get("value", []):
        name = item.get("name", "")
        if not item.get("folder"):
            files.append({
                "name": name,
                "id": item.get("id"),
                "size": item.get("size"),
                "path": item.get("parentReference", {}).get("path", ""),
                "webUrl": item.get("webUrl", "")
            })
    logger.info(f"search_files found {len(files)} files for query '{query}'")
    return files


async def read_file_content(file_id: str, file_name: str) -> dict:
    """קורא את תוכן הקובץ לפי סוג ומחזיר תוכן + לינק"""
    token = await get_token()
    site_id = await get_site_id()

    # קבל מטאדאטה של הקובץ כולל URL
    meta_url = f"{GRAPH_BASE}/sites/{site_id}/drive/items/{file_id}"
    async with aiohttp.ClientSession() as session:
        async with session.get(meta_url, headers={"Authorization": f"Bearer {token}"}) as resp:
            meta = await resp.json()
            web_url = meta.get("webUrl", "")

    # הורד את תוכן הקובץ - עם retry
    content_url = f"{GRAPH_BASE}/sites/{site_id}/drive/items/{file_id}/content"
    logger.info(f"read_file: {file_name} ({file_id})")
    content = None
    for attempt in range(3):
        try:
            async with aiohttp.ClientSession() as session:
                async with session.get(content_url, headers={"Authorization": f"Bearer {token}"}, allow_redirects=True) as resp:
                    logger.info(f"read_file attempt {attempt+1} status: {resp.status}")
                    if resp.status == 200:
                        content = await resp.read()
                        if content and len(content) > 0:
                            break
        except Exception as e:
            logger.error(f"read_file attempt {attempt+1} error: {e}")
        import asyncio as aio
        await aio.sleep(1)

    if not content:
        return {"content": "(שגיאה בהורדת הקובץ — נסה שוב)", "webUrl": web_url}

    lower = file_name.lower()
    text = ""
    if lower.endswith(".pdf"):
        text = _extract_pdf(content)
    elif lower.endswith(".docx"):
        text = _extract_docx(content)
    elif lower.endswith(".doc"):
        text = "(קובץ .doc ישן — לא ניתן לקרוא. יש להמיר ל-.docx)"
    elif lower.endswith(".xlsx") or lower.endswith(".xls"):
        text = _extract_xlsx(content)
    elif lower.endswith(".txt") or lower.endswith(".csv"):
        text = content.decode("utf-8", errors="replace")
    else:
        text = f"(סוג קובץ לא נתמך: {file_name})"

    return {"content": text, "webUrl": web_url}


def _extract_pdf(content: bytes) -> str:
    text = []
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text.append(t)
    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        return f"(שגיאה בקריאת PDF: {e})"
    return "\n".join(text)


def _extract_docx(content: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(content))
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        return f"(שגיאה בקריאת DOCX: {e})"


def _extract_xlsx(content: bytes) -> str:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        lines = []
        full_lines = []
        summary_keywords = ["total", "סה\"כ", "סהכ", "sum", "totals", "grand total"]

        for sheet in wb.worksheets:
            sheet_header = f"\n=== Sheet: {sheet.title} ==="
            lines.append(sheet_header)
            full_lines.append(sheet_header)
            headers = []
            all_rows = list(sheet.iter_rows(values_only=True))
            total_rows = len(all_rows)

            for row_idx, row in enumerate(all_rows, 1):
                cells = [str(c) if c is not None else "" for c in row]
                if row_idx == 1:
                    headers = cells
                    header_line = "כותרות: " + " | ".join(cells)
                    lines.append(header_line)
                    full_lines.append(header_line)
                else:
                    if not any(c.strip() for c in cells):
                        continue
                    first_cell = cells[0].lower().strip()
                    is_summary = any(kw in first_cell for kw in summary_keywords)

                    if headers:
                        parts = []
                        for h, c in zip(headers, cells):
                            if c.strip():
                                col_name = h if h.strip() else "עמודה"
                                parts.append(f"{col_name}: {c}")
                        if parts:
                            prefix = f"⭐ שורה {row_idx} (סיכום)" if is_summary else f"שורה {row_idx}"
                            line = f"{prefix}: " + " | ".join(parts)
                            full_lines.append(line)
                            if is_summary:
                                lines.append(line)
                    else:
                        full_lines.append(" | ".join(cells))

            lines.append(f"(סה\"כ {total_rows} שורות בגיליון {sheet.title})")
            full_lines.append(f"(סה\"כ {total_rows} שורות בגיליון {sheet.title})")

        # אם הטקסט המלא קצר מספיק — שלח הכל
        full_text = "\n".join(full_lines)
        if len(full_text) <= 16000:
            return full_text

        # אחרת שלח סיכום + כותרות מכל גיליון, ואז כמה שאפשר מהנתונים המלאים
        summary_text = "\n".join(lines)
        remaining = 16000 - len(summary_text) - 100
        if remaining > 0:
            return summary_text + "\n\n--- נתונים מפורטים (חלקי) ---\n" + full_text[:remaining]
        return summary_text

    except Exception as e:
        logger.error(f"XLSX extraction error: {e}")
        return f"(שגיאה בקריאת XLSX: {e})"